"""XlsxExporter：生成 .xlsx 评分清单（openpyxl）。

列顺序与 CSV（MANIFEST_COLUMNS）完全一致，分数列带红→黄→绿条件格式，
表头蓝底白字粗体，并开启自动筛选。本模块**只在 entry.py 的 --xlsx 分支内被
lazy import**，因此基础安装（未装 openpyxl）不受影响，CLI/MCP 仍可正常启动。
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from flashgaze.models.schemas import (
    MANIFEST_COLUMNS,  # 列顺序单一来源（C-09 迁至 schemas.py）
    ImageScore,
)
from flashgaze.utils.atomic_write import atomic_write_bytes
from flashgaze.utils.logging import get_logger

logger = get_logger()

HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
HEADER_FONT = Font(color="FFFFFF", bold=True)
# 分数列在 MANIFEST_COLUMNS 中的索引（0-based）：file,path,sharpness,exposure,
# composition,vlm_adjustment,score -> 第 6 个 -> 列 G
_SCORE_COLUMN_INDEX = MANIFEST_COLUMNS.index("score")  # 6


class XlsxExporter:
    """生成 .xlsx 评分清单。列顺序与 CSV 完全一致，分数列带红黄绿条件格式。"""

    def write(self, scored: list[ImageScore], path: str) -> None:
        if not scored:
            logger.warning("XlsxExporter.write 收到空列表，仍生成空表头文件")
        wb = Workbook()
        ws = wb.active
        ws.title = "FlashGaze 评分"
        ws.append(list(MANIFEST_COLUMNS))
        self._style_header(ws)
        for s in scored:
            ws.append(self._row(s))
        self._apply_score_color_scale(ws, len(scored))
        last_col = get_column_letter(len(MANIFEST_COLUMNS))
        ws.auto_filter.ref = f"A1:{last_col}{len(scored) + 1}"
        # 原子写：openpyxl 直写路径不可原子，改用 BytesIO + atomic_write_bytes
        buf = io.BytesIO()
        wb.save(buf)
        atomic_write_bytes(path, buf.getvalue())

    @staticmethod
    def _style_header(ws: Worksheet) -> None:
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

    @staticmethod
    def _row(s: ImageScore) -> list[object]:
        # 复用 ImageScore.to_row()（C-09 单一来源）：与 CSV 完全一致，
        # None 维度写空串，selected/deduped 保留原生布尔（Excel 布尔）
        return s.to_row()

    @staticmethod
    def _apply_score_color_scale(ws: Worksheet, n_rows: int) -> None:
        if n_rows <= 0:
            return
        col = get_column_letter(_SCORE_COLUMN_INDEX + 1)  # 1-based
        rng = f"{col}2:{col}{n_rows + 1}"
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="FF0000",
                mid_type="num",
                mid_value=60,
                mid_color="FFFF00",
                end_type="num",
                end_value=100,
                end_color="00FF00",
            ),
        )
